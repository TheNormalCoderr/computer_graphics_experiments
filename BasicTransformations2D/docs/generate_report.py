"""
generate_report.py — build BasicTransformations2D_Experiment_Report.docx
Run from repo root with venv active:
  python BasicTransformations2D/docs/generate_report.py
"""

import os
import sys
import math

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT  = os.path.abspath(os.path.join(SCRIPT_DIR, "..", ".."))
OUTPUT_DIR = os.path.join(REPO_ROOT, "BasicTransformations2D", "outputs")
DOCS_DIR   = os.path.join(REPO_ROOT, "BasicTransformations2D", "docs")
DOCX_OUT   = os.path.join(DOCS_DIR, "BasicTransformations2D_Experiment_Report.docx")

SRC_DIR = os.path.join(REPO_ROOT, "BasicTransformations2D", "src")
sys.path.insert(0, SRC_DIR)
from basic_transformations import TEST_CASES, apply_transformation, format_matrix  # type:ignore


def main():
    try:
        from docx import Document  # type:ignore
        from docx.shared import Pt, Inches, RGBColor  # type:ignore
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type:ignore
    except ImportError:
        print("ERROR: python-docx not installed. Run: pip install python-docx")
        sys.exit(1)

    doc = Document()

    # 1 inch margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)

    def set_font(para, size=11, bold=False, italic=False, align=None, color=None):
        if align:
            para.alignment = align
        for run in para.runs:
            run.font.name   = "Times New Roman"
            run.font.size   = Pt(size)
            run.font.bold   = bold
            run.font.italic = italic
            if color:
                run.font.color.rgb = RGBColor(*color)

    def add_para(text="", bold=False, italic=False, size=11, align=None, space_before=0, space_after=6):
        p = doc.add_paragraph(text)
        if align:
            p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(space_after)
        for run in p.runs:
            run.font.name   = "Times New Roman"
            run.font.size   = Pt(size)
            run.font.bold   = bold
            run.font.italic = italic
        return p

    def add_run(para, text, bold=False, italic=False, size=11, color=None):
        run = para.add_run(text)
        run.font.name   = "Times New Roman"
        run.font.size   = Pt(size)
        run.font.bold   = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = RGBColor(*color)
        return run

    def section_label(label):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(label)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True
        return p

    def sub_label(label):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after  = Pt(2)
        r = p.add_run(label)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        return p

    def body(text, size=11, space_before=0, space_after=6):
        p = doc.add_paragraph(text)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(space_after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
        return p

    def code_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        for index, line in enumerate(text.strip("\n").splitlines()):
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            if index < len(text.strip("\n").splitlines()) - 1:
                run.add_break()
        return p

    def bullet(text, level=0):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        return p

    def caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after  = Pt(10)
        r = p.add_run(text)
        r.font.name   = "Times New Roman"
        r.font.size   = Pt(10)
        r.font.italic = True
        return p

    def add_image(path, width=Inches(5.0)):
        if not os.path.exists(path):
            print(f"[Warning] Image not found: {path}")
            return None
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(2)
        run = p.add_run()
        run.add_picture(path, width=width)
        return p

    def add_table(headers, rows):
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hrow = t.rows[0]
        for i, h in enumerate(headers):
            cell = hrow.cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"
            cell.paragraphs[0].runs[0].font.size = Pt(10)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ri, row in enumerate(rows):
            drow = t.rows[ri + 1]
            for ci, val in enumerate(row):
                cell = drow.cells[ci]
                cell.text = str(val)
                cell.paragraphs[0].runs[0].font.name = "Times New Roman"
                cell.paragraphs[0].runs[0].font.size = Pt(10)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return t

    # ══════════════════════════════════════════════════════════════════════
    # 1. Centered experiment number & title
    # ══════════════════════════════════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Experiment 6")
    r.font.name = "Times New Roman"; r.font.size = Pt(14); r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(14)
    r = p.add_run("Implementation of 2D Basic Transformations using OpenGL in Python")
    r.font.name = "Times New Roman"; r.font.size = Pt(14); r.font.bold = True

    # ══════════════════════════════════════════════════════════════════════
    # 2. Aim
    # ══════════════════════════════════════════════════════════════════════
    section_label("Aim:")
    body(
        "To implement the 2D basic geometric transformations (Translation, Scaling, and Rotation) "
        "in Python using homogeneous coordinates and PyOpenGL, and to render the visual comparison "
        "between the original and transformed polygon on an annotated Cartesian grid."
    )

    # ══════════════════════════════════════════════════════════════════════
    # 3. Problem Statement
    # ══════════════════════════════════════════════════════════════════════
    section_label("Problem Statement:")
    body(
        "Given an n-sided 2D polygon defined by vertices V = {(x1, y1), (x2, y2), ..., (xn, yn)}, "
        "apply a specified basic geometric transformation with parameters (tx, ty for translation; "
        "sx, sy for scaling; and angle theta for rotation about the origin) to obtain the transformed "
        "vertex set V'. The transformation must be formulated using 3x3 homogeneous matrix representation, "
        "and both the original and transformed polygons must be displayed simultaneously on an interactive "
        "reference coordinate system."
    )
    sub_label("Input / Setup:")
    body(
        "A list of 2D coordinates representing polygon vertices; transformation type and associated numerical "
        "parameters; canvas resolution 960 x 700 pixels with a 20 x 16 reference grid."
    )
    sub_label("Process / Constraint:")
    body(
        "Compute the 3x3 homogeneous transformation matrix M; map each vertex [x, y, 1]^T to [x', y', 1]^T = M · [x, y, 1]^T; "
        "render original shape in blue and transformed shape in red."
    )
    sub_label("Output:")
    body("Rendered comparison windows and exported PNG images demonstrating each transformation.")

    # ══════════════════════════════════════════════════════════════════════
    # 4. Theory
    # ══════════════════════════════════════════════════════════════════════
    section_label("Theory:")
    body(
        "Geometric transformations are fundamental operations in computer graphics used to reposition, "
        "resize, and orient objects within a scene. By using Cartesian coordinates alone, translation involves "
        "vector addition whereas scaling and rotation involve matrix multiplication. To unify all affine transformations "
        "into a single consistent matrix multiplication operation, homogeneous coordinates are used."
    )
    body(
        "In homogeneous coordinates, a 2D point (x, y) is represented as a 3D column vector [x, y, 1]^T. "
        "The standard 3x3 transformation matrices are defined as follows:"
    )
    bullet("Translation: Displaces a point by (tx, ty): x' = x + tx, y' = y + ty.")
    bullet("Scaling: Multiplies coordinate distances from the origin by factors sx and sy: x' = x · sx, y' = y · sy.")
    bullet("Rotation: Rotates a point counter-clockwise about the origin by angle θ: x' = x cos θ - y sin θ, y' = x sin θ + y cos θ.")

    # ══════════════════════════════════════════════════════════════════════
    # 5. Algorithm / Methodology
    # ══════════════════════════════════════════════════════════════════════
    section_label("Algorithm / Methodology:")
    body("1. Input the polygon vertices V = {(x1, y1), ..., (xn, yn)} and transformation parameters.")
    body("2. Construct the 3x3 transformation matrix M according to the chosen transformation:")
    body("   - Translation: M = [[1, 0, tx], [0, 1, ty], [0, 0, 1]]")
    body("   - Scaling: M = [[sx, 0, 0], [0, sy, 0], [0, 0, 1]]")
    body("   - Rotation: M = [[cos θ, -sin θ, 0], [sin θ, cos θ, 0], [0, 0, 1]]")
    body("3. For each vertex (x, y) in V, compute [x', y', 1]^T = M · [x, y, 1]^T.")
    body("4. Extract the transformed Cartesian coordinates (x', y').")
    body("5. Render the Cartesian grid, axes, original shape (blue), and transformed shape (red) using OpenGL.")

    # ══════════════════════════════════════════════════════════════════════
    # 6. Implementation
    # ══════════════════════════════════════════════════════════════════════
    section_label("Implementation:")

    body(
        "The implementation represents every 2D point in homogeneous form and applies one 3x3 matrix for the selected "
        "operation. Translation changes the final column, scaling changes the diagonal entries, and rotation is built "
        "from the sine and cosine of the angle."
    )
    sub_label("Translation, Scaling, and Rotation Code:")
    code_block("""
def translation_matrix(tx, ty):
    return [[1.0, 0.0, float(tx)],
            [0.0, 1.0, float(ty)],
            [0.0, 0.0, 1.0]]

def scaling_matrix(sx, sy):
    return [[float(sx), 0.0, 0.0],
            [0.0, float(sy), 0.0],
            [0.0, 0.0, 1.0]]

def rotation_matrix(angle_deg):
    theta = math.radians(angle_deg)
    c, s = math.cos(theta), math.sin(theta)
    return [[c, -s, 0.0],
            [s, c, 0.0],
            [0.0, 0.0, 1.0]]
""")

    sub_label("Applying the Selected Transformation:")
    code_block("""
def apply_transformation(vertices, transform_name, params):
    if transform_name == "translation":
        M = translation_matrix(params["tx"], params["ty"])
    elif transform_name == "scaling":
        M = scaling_matrix(params["sx"], params["sy"])
    elif transform_name == "rotation":
        M = rotation_matrix(params["angle_deg"])
    else:
        raise ValueError(f"Unknown transformation: {transform_name}")

    transformed = [mat3_transform_point(M, x, y) for x, y in vertices]
    return transformed, M
""")

    # 7. Test Cases
    # ══════════════════════════════════════════════════════════════════════
    section_label("Test Cases:")
    tc_table_headers = ["Test Case", "Original Vertices", "Transformation", "Parameters", "Transformed Vertices"]
    tc_table_rows = []

    for tc in TEST_CASES:
        transformed, M = apply_transformation(tc["vertices"], tc["transform"], tc["params"])
        t_rounded = [(round(x, 2), round(y, 2)) for x, y in transformed]
        param_str = ", ".join(f"{k}={v}" for k, v in tc["params"].items())
        tc_table_rows.append([
            tc["label"].split("—")[0].strip(),
            str(tc["vertices"]),
            tc["transform"].capitalize(),
            param_str,
            str(t_rounded),
        ])

    add_table(tc_table_headers, tc_table_rows)
    caption("Table 6.1: Test Cases for 2D Basic Transformations")

    # Images
    for tc in TEST_CASES:
        sub_label(tc["label"])
        img_path = os.path.join(OUTPUT_DIR, tc["file"])
        add_image(img_path, width=Inches(5.5))
        caption(f"Visual result for {tc['label']}")

    # ══════════════════════════════════════════════════════════════════════
    # 7. Result
    # ══════════════════════════════════════════════════════════════════════
    section_label("Result:")
    body(
        "The 2D basic transformation algorithms for translation, scaling, and rotation were successfully implemented "
        "using homogeneous coordinate matrices. All test cases produced mathematically accurate transformed vertices "
        "and visually verified plots rendered on the OpenGL canvas."
    )

    # ══════════════════════════════════════════════════════════════════════
    # 8. Conclusion
    # ══════════════════════════════════════════════════════════════════════
    section_label("Conclusion:")
    body(
        "Homogeneous coordinate representation enables all basic 2D geometric transformations to be treated uniformly "
        "as 3x3 matrix multiplications. This provides a clean mathematical foundation for the graphics pipeline, allowing "
        "seamless compounding and consistent hardware or software implementation."
    )

    # ══════════════════════════════════════════════════════════════════════
    # 9. Source Code and Analysis (GitHub)
    # ══════════════════════════════════════════════════════════════════════
    section_label("Source Code and Analysis (GitHub):")
    body("https://github.com/TheNormalCoderr/computer_graphics_experiments")

    os.makedirs(DOCS_DIR, exist_ok=True)
    doc.save(DOCX_OUT)
    print(f"[Success] Report generated: {DOCX_OUT}")


if __name__ == "__main__":
    main()
